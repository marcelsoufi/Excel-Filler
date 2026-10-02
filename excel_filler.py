# -*- coding: utf-8 -*-
"""
Excel Report Filler  —  manual cell-mapping edition
===================================================
Run:           python excel_report_filler.py
Requirements:  pip install PySide6 openpyxl

How it works
  * Every Excel file is handled independently (own preview, own mappings, own state).
  * You pick cells yourself: click / drag in the sheet preview, or type addresses
    (e.g.  Survey!B5   |   'Sheet 2'!C3:C40 ; Info!B2).
  * Sources may come from several sheets AND from other files (button «Source files»).
    Sheets of other files are written like  '[Other.xlsx]Survey'!B5.
  * Every source cell is paired with exactly one target cell (same order). You can choose
    row-by-row or column-by-column order and check the pairs in a preview before adding.
    Several source cells -> ONE target cell is also possible (first non-empty / join / sum).
  * The ordered list of mappings can be saved as a named TEMPLATE and reused on other files;
    if sheet names differ, the app asks which sheet to use for each one.
  * On «Execute» you can apply the same operation to all open files, to some of them,
    or only to the current one. Results are always saved as NEW copies beside the originals.
  * Light / dark theme switch (remembered).
"""
from __future__ import annotations

import datetime as _dt
import json
import re
import sys
import traceback
import uuid
from collections import Counter
from copy import deepcopy
from pathlib import Path
from string import Template
from typing import Any, Callable, Optional

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

APP_NAME = "Excel Report Filler"
CONFIG_DIR = Path.home() / ".excel_report_filler"
_OLD_CONFIG_DIR = Path.home() / ".iga_visit_filler"          # earlier name: templates are migrated once
if not CONFIG_DIR.exists() and (_OLD_CONFIG_DIR / "templates.json").exists():
    try:
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        (CONFIG_DIR / "templates.json").write_bytes((_OLD_CONFIG_DIR / "templates.json").read_bytes())
    except Exception:
        pass
TEMPLATES_FILE = CONFIG_DIR / "templates.json"
SETTINGS_FILE = CONFIG_DIR / "settings.json"
MAX_ROWS, MAX_COLS, MAX_CELLS = 3000, 80, 5000
DEV_EMAIL = "yousefsoufi2@gmail.com"
ICON_FILE = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "excel_report_filler.ico"

MODES = {"first": "First non-empty value", "join": "Join values (one per line)", "sum": "Sum of numeric values"}


class _Cancelled(Exception):
    pass


def load_settings() -> dict:
    try:
        return json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_settings(d: dict):
    try:
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        SETTINGS_FILE.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
    except Exception:
        pass


# ════════════════════════════════════════════════════════════════════════════
#  1) Values
# ════════════════════════════════════════════════════════════════════════════
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")


def is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.replace("\xa0", " ").strip() == "")


def to_number(v: Any) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.translate(_DIGITS).replace(",", "").replace("٫", ".").strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s) if "." in s else int(s)
    return None


def convert_scalar(v: Any, convert_digits: bool) -> Any:
    if is_blank(v):
        return None
    if isinstance(v, str):
        s = v.replace("\xa0", " ").strip()
        if not convert_digits:
            return s
        s = s.translate(_DIGITS)
        if re.fullmatch(r"\d+(\.\d+)?", s) and not (len(s) > 1 and s.startswith("0") and "." not in s):
            return float(s) if "." in s else int(s)
        return s
    return v


def fmt_display(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ════════════════════════════════════════════════════════════════════════════
#  2) Cell references  (Sheet!B5, 'My Sheet'!B5:C9, '[Other.xlsx]Survey'!B5 ; Info!A1)
# ════════════════════════════════════════════════════════════════════════════
_SPLIT = re.compile(r"[;\n،,](?=(?:[^']*'[^']*')*[^']*$)")
_CELL = re.compile(r"\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?")


def qsheet(name: str) -> str:
    return name if re.fullmatch(r"\w+", name) else "'" + name.replace("'", "''") + "'"


def fmt_ref(sheet: str, coord: str) -> str:
    return f"{qsheet(sheet)}!{coord}"


def key_label(key: str) -> str:
    """Human label for a sheet key: external sheets are shown as  file.xlsx › Sheet."""
    if key.startswith("[") and "]" in key:
        lab, name = key[1:].split("]", 1)
        return f"{lab} › {name}"
    return key


def parse_refs(text: str, default_sheet: str, by_col: bool = False) -> list[tuple[str, str]]:
    """Expands the text into an ordered list of (sheet, 'B5'). Ranges are expanded row by row,
    or column by column when by_col is True."""
    cells: list[tuple[str, str]] = []
    for part in _SPLIT.split(text or ""):
        part = part.strip()
        if not part:
            continue
        sheet = default_sheet
        if "!" in part:
            sheet, part = part.rsplit("!", 1)
            sheet, part = sheet.strip(), part.strip()
            if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
                sheet = sheet[1:-1].replace("''", "'")
        m = _CELL.fullmatch(part)
        if not m:
            raise ValueError(f"Invalid cell reference: «{part}»")
        try:
            c1 = column_index_from_string(m[1].upper())
            c2 = column_index_from_string(m[3].upper()) if m[3] else c1
        except ValueError:
            raise ValueError(f"Invalid column in «{part}»")
        r1 = int(m[2])
        r2 = int(m[4]) if m[4] else r1
        r1, r2 = sorted((r1, r2))
        c1, c2 = sorted((c1, c2))
        if r1 < 1:
            raise ValueError(f"Invalid row in «{part}»")
        rows, cols = range(r1, r2 + 1), range(c1, c2 + 1)
        order = ((r, c) for c in cols for r in rows) if by_col else ((r, c) for r in rows for c in cols)
        for r, c in order:
            cells.append((sheet, f"{get_column_letter(c)}{r}"))
            if len(cells) > MAX_CELLS:
                raise ValueError(f"Too many cells (max {MAX_CELLS}).")
    return cells


def split_coord(coord: str) -> tuple[int, int]:
    col, row = coordinate_from_string(coord)
    return row, column_index_from_string(col)


# ════════════════════════════════════════════════════════════════════════════
#  3) Mappings (rules), templates, sheet matching
# ════════════════════════════════════════════════════════════════════════════
def new_rule(src: list, dst: list, mode: str = "first", enabled: bool = True) -> dict:
    return {"id": uuid.uuid4().hex[:8], "src": [list(s) for s in src], "dst": list(dst),
            "mode": mode if mode in MODES else "first", "enabled": bool(enabled)}


class Templates:
    """Named lists of mappings, stored in ~/.excel_report_filler/templates.json (cell addresses only, no data)."""

    def __init__(self, path: Path = TEMPLATES_FILE):
        self.path = path
        self.data: dict[str, dict] = {}
        try:
            if path.exists():
                self.data = json.loads(path.read_text(encoding="utf-8")).get("templates", {})
        except Exception:
            self.data = {}

    def names(self) -> list[str]:
        return sorted(self.data, key=str.lower)

    def get(self, name: str) -> list[dict]:
        out = []
        for r in deepcopy(self.data.get(name, {}).get("rules", [])):
            try:
                out.append(new_rule(r["src"], r["dst"], r.get("mode", "first"), r.get("enabled", True)))
            except Exception:
                continue
        return out

    def put(self, name: str, rules: list[dict], index: Optional[dict] = None, ext_files: Optional[dict] = None,
            seq: Optional[dict] = None):
        old = self.data.get(name, {})
        self.data[name] = {
            "updated": _dt.datetime.now().isoformat(timespec="seconds"),
            "rules": [{"src": r["src"], "dst": r["dst"], "mode": r["mode"], "enabled": r["enabled"]} for r in rules],
            "index": index if index is not None else old.get("index", {}),      # sheet position in the workbook
            "aliases": old.get("aliases", {}),                                  # last sheet chosen for each template sheet
            "ext_files": ext_files if ext_files is not None else old.get("ext_files", {}),   # extra source files: name -> path
            "seq": seq if seq is not None else old.get("seq", {}),                          # sequential-fill settings
        }
        self.save()

    def info(self, name: str) -> tuple[dict, dict]:
        d = self.data.get(name, {})
        return d.get("index", {}), d.get("aliases", {})

    def get_seq(self, name: str) -> dict:
        return dict(self.data.get(name, {}).get("seq", {}))

    def ext_files(self, name: str) -> dict:
        return dict(self.data.get(name, {}).get("ext_files", {}))

    def set_ext_file(self, name: str, label: str, path: str):
        if name in self.data:
            self.data[name].setdefault("ext_files", {})[label] = path
            self.save()

    def set_aliases(self, name: str, aliases: dict):
        if name in self.data:
            self.data[name]["aliases"] = aliases
            self.save()

    def delete(self, name: str):
        if self.data.pop(name, None) is not None:
            self.save()

    def save(self):
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps({"templates": self.data}, ensure_ascii=False, indent=1), encoding="utf-8")


def _base(key: str) -> str:
    return key.split("]", 1)[1] if key.startswith("[") and "]" in key else key


def suggest_sheet(tpl_sheet: str, file_sheets: list[str], index: Optional[int] = None,
                  alias: Optional[str] = None, main_count: Optional[int] = None) -> str:
    """Best guess for which sheet of the current file plays the role of a template sheet ('' = no guess).
    file_sheets lists the file's own sheets first, then sheets of attached source files."""
    if tpl_sheet in file_sheets:
        return tpl_sheet
    if alias and alias in file_sheets:                   # what the user chose last time for this template
        return alias
    ext = tpl_sheet.startswith("[")
    pool = [f for f in file_sheets if f.startswith("[") == ext]
    norm = lambda x: re.sub(r"[\W_]+", "", _base(x)).lower()
    nt = norm(tpl_sheet)
    same = [f for f in pool if norm(f) == nt]
    if len(same) == 1:
        return same[0]
    if len(nt) >= 3:
        part = [f for f in pool if len(norm(f)) >= 3 and (nt in norm(f) or norm(f) in nt)]
        if len(part) == 1:
            return part[0]
    n_main = main_count if main_count is not None else len(file_sheets)
    if not ext and index is not None and 0 <= index < n_main:     # same position as when the template was saved
        return file_sheets[index]
    return ""


class Book:
    """Read-only view (calculated values) of the main workbook plus any attached source files."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.ext: dict[str, Any] = {}              # label (file name) -> workbook
        self.ext_paths: dict[str, Path] = {}
        self._anchor: dict[tuple[str, str], str] = {}

    @staticmethod
    def split_key(key: str) -> tuple[str, str]:
        if key.startswith("[") and "]" in key:
            lab, name = key[1:].split("]", 1)
            return lab, name
        return "", key

    def main_sheets(self) -> list[str]:
        return list(self.wb.sheetnames)

    def all_keys(self) -> list[str]:
        return self.main_sheets() + [f"[{lab}]{n}" for lab, w in self.ext.items() for n in w.sheetnames]

    def is_main(self, key: str) -> bool:
        return self.split_key(key)[0] == "" and key in self.wb.sheetnames

    def has(self, key: str) -> bool:
        lab, name = self.split_key(key)
        if lab:
            return lab in self.ext and name in self.ext[lab].sheetnames
        return name in self.wb.sheetnames

    def ws(self, key: str):
        lab, name = self.split_key(key)
        return (self.ext[lab] if lab else self.wb)[name]

    def add_external(self, path: Path) -> str:
        path = Path(path)
        lab, i = path.name, 2
        while lab in self.ext:
            lab = f"{path.stem} ({i}){path.suffix}"
            i += 1
        self.ext[lab] = openpyxl.load_workbook(path, data_only=True)
        self.ext_paths[lab] = path
        return lab

    def remove_external(self, lab: str):
        self.ext.pop(lab, None)
        self.ext_paths.pop(lab, None)
        self._anchor.clear()

    def anchor(self, sheet: str, coord: str) -> str:
        key = (sheet, coord)
        if key not in self._anchor:
            a = coord
            for rng in self.ws(sheet).merged_cells.ranges:
                if coord in rng:
                    a = rng.start_cell.coordinate
                    break
            self._anchor[key] = a
        return self._anchor[key]

    def get(self, sheet: str, coord: str) -> Any:
        return self.ws(sheet)[self.anchor(sheet, coord)].value


def rule_sheets_ok(book: Book, rule: dict) -> bool:
    return book.is_main(rule["dst"][0]) and all(book.has(s[0]) for s in rule["src"])


def rule_value(book: Book, rule: dict, override: Optional[str], convert_digits: bool) -> Any:
    if override is not None:
        return convert_scalar(override, convert_digits)
    if not rule_sheets_ok(book, rule):
        return None
    vals = [v for v in (book.get(sh, co) for sh, co in rule["src"]) if not is_blank(v)]
    if not vals:
        return None
    mode = rule["mode"]
    if mode == "sum":
        nums = [n for n in (to_number(v) for v in vals) if n is not None]
        return sum(nums) if nums else None
    if mode == "join":
        return "\n".join(fmt_display(convert_scalar(v, convert_digits)) for v in vals)
    return convert_scalar(vals[0], convert_digits)


def rule_status(book: Book, r: dict, value: Any, edited: bool, dups: Counter, overwrite: bool) -> tuple[str, bool]:
    """(status text, will be written?)"""
    if not r["enabled"]:
        return "Off", False
    if not rule_sheets_ok(book, r):
        return "Sheet missing", False
    if value is None:
        return "Source empty", False
    existing = book.get(*r["dst"])
    if not is_blank(existing) and not overwrite:
        return "Target filled", False
    if dups[tuple(r["dst"])] > 1:
        return "Duplicate target", True
    if not is_blank(existing):
        return "Overwrites", True
    return ("Edited" if edited else "Ready"), True


def plan_writes(book: Book, rules: list[dict], overwrite: bool, convert_digits: bool,
                overrides: Optional[dict] = None) -> list[tuple[str, str, Any]]:
    overrides = overrides or {}
    dups = Counter(tuple(r["dst"]) for r in rules if r["enabled"])
    out = []
    for r in rules:
        v = rule_value(book, r, overrides.get(r["id"]), convert_digits)
        if rule_status(book, r, v, r["id"] in overrides, dups, overwrite)[1]:
            out.append((r["dst"][0], r["dst"][1], v))
    return out


def shift_coord(coord: str, dr: int, dc: int) -> str:
    r, c = split_coord(coord)
    r2, c2 = r + dr, c + dc
    return "" if r2 < 1 or c2 < 1 else f"{get_column_letter(c2)}{r2}"


def _resolve_one(sh: str, ref_sheets: list[str], rec_sheets: list[str], override: Optional[str] = None) -> str:
    """Sheet of a record file that plays the role of sample sheet `sh`: the user's choice, else same name,
    else same position ('' = none)."""
    if override and override in rec_sheets:
        return override
    if sh in rec_sheets:
        return sh
    if sh in ref_sheets and ref_sheets.index(sh) < len(rec_sheets):
        return rec_sheets[ref_sheets.index(sh)]
    return ""


def seq_resolve_sheets(ref_sheets: list[str], rec_sheets: list[str], used: list[str],
                       overrides: Optional[dict] = None) -> dict[str, str]:
    m = {}
    for sh in used:
        r = _resolve_one(sh, ref_sheets, rec_sheets, (overrides or {}).get(sh))
        if not r:
            raise ValueError(f"sheet «{sh}» not found")
        m[sh] = r
    return m


def build_sequence(main: Book, rules: list[dict], ref: str, files: list[Path], direction: str, step: int,
                   auto_start: bool, counter: Optional[tuple[str, str, int]], overwrite: bool, convert_digits: bool,
                   progress: Optional[Callable[[int, int, str], None]] = None,
                   sheet_overrides: Optional[dict] = None) -> dict:
    """Sequential fill. The mappings were made on a SAMPLE source file (attached under label `ref`) and point at the
    FIRST record of the target table. The same source cells are read from every file in `files`; file number i is
    written `i * step` rows down (or columns right) from the first record. Rules that don't read from the sample file
    are 'fixed' and written once."""
    prefix = f"[{ref}]"
    active = [r for r in rules if r["enabled"]]
    seq_rules = [r for r in active if any(s[0].startswith(prefix) for s in r["src"])]
    ids = {r["id"] for r in seq_rules}
    fixed = [r for r in active if r["id"] not in ids]
    unit = (step, 0) if direction == "down" else (0, step)
    used = list(dict.fromkeys(Book.split_key(s[0])[1] for r in seq_rules for s in r["src"] if s[0].startswith(prefix)))
    ref_sheets = main.ext[ref].sheetnames
    at = lambda sh, co, k: main.get(sh, shift_coord(co, unit[0] * k, unit[1] * k))

    offset = 0
    if auto_start:                                   # first position where the whole record is still empty
        for k in range(200000):
            if all(is_blank(at(r["dst"][0], r["dst"][1], k)) for r in seq_rules):
                offset = k
                break

    records, writes, i = [], [], 0
    warnings: dict[str, str] = {}
    saved = main.ext[ref]
    try:
        for n, path in enumerate(files):
            if progress:
                progress(n + 1, len(files), path.name)
            try:
                wb = openpyxl.load_workbook(path, data_only=True)
                ov = (sheet_overrides or {}).get(str(path), {})
                smap = seq_resolve_sheets(ref_sheets, wb.sheetnames, used, ov)
            except Exception as ex:
                records.append((path, f"skipped: {ex}", [], None))
                continue
            k = offset + i
            rk = []
            for r in seq_rules:
                nr = deepcopy(r)
                nr["dst"] = [r["dst"][0], shift_coord(r["dst"][1], unit[0] * k, unit[1] * k)]
                nr["src"] = [[f"{prefix}{smap[Book.split_key(s[0])[1]]}", s[1]] if s[0].startswith(prefix) else s
                             for s in r["src"]]
                rk.append(nr)
            main.ext[ref] = wb
            main._anchor.clear()
            w = plan_writes(main, rk, overwrite, convert_digits)
            main.ext[ref] = saved
            main._anchor.clear()
            if not w:
                records.append((path, "skipped: nothing to write (empty source cells or target already filled)", [], None))
                continue
            notes = []
            moved = {a: b for a, b in smap.items() if a != b and ov.get(a) != b}      # chosen by the user = confirmed
            if moved:                                       # matched by position, not by name: worth a look
                notes.append("sheet " + ", ".join(f"«{a}»→«{b}»" for a, b in moved.items()))
            if len(w) < len(seq_rules):
                notes.append(f"only {len(w)} of {len(seq_rules)} cells have values")
            if notes:
                warnings[str(path)] = "; ".join(notes)
            if counter:
                num = (offset + 1 + i) if counter[2] == 0 else counter[2] + i
                w.append((counter[0], shift_coord(counter[1], unit[0] * k, unit[1] * k), num))
            records.append((path, "ok", w, k))
            writes += w
            i += 1
    finally:
        main.ext[ref] = saved
        main._anchor.clear()

    columns = []
    for r in seq_rules:
        sh, co = r["dst"]
        hc = shift_coord(co, -1, 0) if direction == "down" else shift_coord(co, 0, -1)
        hv = main.get(sh, hc) if hc else None
        columns.append((sh, co, hv.strip() if isinstance(hv, str) else ""))
    fixed_writes = plan_writes(main, fixed, overwrite, convert_digits)
    return {"offset": offset, "records": records, "writes": fixed_writes + writes, "columns": columns,
            "counter": counter, "unit": unit, "n_ok": i, "n_fixed": len(fixed_writes), "warnings": warnings,
            "used_sheets": used, "ref_sheets": ref_sheets}


def unique_output_path(src: Path) -> Path:
    out = src.with_name(f"{src.stem}_filled{src.suffix}")
    i = 2
    while out.exists():
        out = src.with_name(f"{src.stem}_filled_{i}{src.suffix}")
        i += 1
    return out


def write_output(src: Path, writes: list[tuple[str, str, Any]], out: Optional[Path] = None) -> tuple[Path, int]:
    wb = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    n = 0
    for sheet, coord, v in writes:
        ws = wb[sheet]
        cell = ws[coord]
        if isinstance(cell, MergedCell):
            for rng in ws.merged_cells.ranges:
                if coord in rng:
                    cell = ws[rng.start_cell.coordinate]
                    break
            else:
                continue
        cell.value = v
        if isinstance(v, (_dt.datetime, _dt.date)) and cell.number_format == "General":
            cell.number_format = "yyyy-mm-dd"
        if isinstance(v, str) and "\n" in v:
            al = cell.alignment
            cell.alignment = Alignment(wrap_text=True, horizontal=al.horizontal, vertical=al.vertical or "top")
        n += 1
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    out = out or unique_output_path(src)
    wb.save(out)
    return out, n


# ════════════════════════════════════════════════════════════════════════════
#  4) Themes
# ════════════════════════════════════════════════════════════════════════════
THEMES = {
    "light": dict(bg="#F3F5F9", card="#FFFFFF", text="#1F2A44", muted="#6B7690", border="#E3E7EE", border2="#CBD3E1",
                  hover="#EEF3FF", hover_b="#9DB6F5", primary="#2F6FED", primary_h="#235BD1", success="#1E9E6A",
                  success_h="#178356", disabled_t="#A3ACBF", disabled_bg="#E3E7EE", grid="#EEF1F6", sel="#C9DBFF",
                  sel_list="#E4ECFF", header="#F3F5F9", tab="#EEF1F6",
                  src="#CFE0FF", dst="#CDEFD9", bad="#FDE2E2", warn="#FFF3D6", ok="#DDF5E8"),
    "dark": dict(bg="#12151C", card="#1B1F2A", text="#E4E8F1", muted="#8C96AB", border="#2A3040", border2="#39415A",
                 hover="#242B3B", hover_b="#4C6FD0", primary="#3D7BFF", primary_h="#5A90FF", success="#23A874",
                 success_h="#2CBE85", disabled_t="#5F687C", disabled_bg="#232838", grid="#252B3A", sel="#2B4A85",
                 sel_list="#24365E", header="#222838", tab="#232838",
                 src="#2B4B8C", dst="#1F5C41", bad="#5A2A2E", warn="#5A4A22", ok="#1F4A38"),
}
THEME_NAME = "light"
THEME = THEMES["light"]

QSS = Template("""
* { font-family: 'Segoe UI', 'Tahoma', 'Arial', 'DejaVu Sans'; font-size: 10.5pt; color: $text; }
QToolTip { color: $text; background: $card; border: 1px solid $border2; }
QMenu { background: $card; color: $text; border: 1px solid $border2; }
QMenu::item { padding: 5px 22px; }
QMenu::item:selected { background: $sel_list; }
QMenu::item:disabled { color: $disabled_t; }
QMessageBox, QDialog { background: $bg; }
QCheckBox, QLabel { background: transparent; }
QFrame#card { background: $card; border: 1px solid $border; border-radius: 12px; }
QLabel#title { font-size: 15pt; font-weight: 700; }
QLabel#muted { color: $muted; }
QLabel#h { font-weight: 600; }
QPushButton { background: $card; border: 1px solid $border2; border-radius: 8px; padding: 6px 12px; }
QPushButton:hover { background: $hover; border-color: $hover_b; }
QPushButton:disabled { color: $disabled_t; background: $bg; }
QPushButton:checked { background: $sel_list; border-color: $primary; font-weight: 600; }
QPushButton::menu-indicator { width: 0px; }
QPushButton#primary { background: $primary; border: 1px solid $primary; color: white; font-weight: 600; }
QPushButton#primary:hover { background: $primary_h; }
QPushButton#success { background: $success; border: 1px solid $success; color: white; font-weight: 600; padding: 9px 14px; }
QPushButton#success:hover { background: $success_h; }
QPushButton#primary:disabled, QPushButton#success:disabled { color: $disabled_t; background: $disabled_bg; border-color: $disabled_bg; }
QPushButton#sig { border: none; background: transparent; color: $muted; padding: 2px 8px; font-size: 9pt; }
QPushButton#sig:hover { color: $primary; background: transparent; }
QComboBox, QLineEdit { background: $card; border: 1px solid $border2; border-radius: 8px; padding: 5px 8px; }
QComboBox QAbstractItemView { background: $card; color: $text; selection-background-color: $sel_list; selection-color: $text; }
QListWidget { background: $card; border: 1px solid $border; border-radius: 10px; padding: 4px; }
QListWidget::item { padding: 6px; border-radius: 6px; }
QListWidget::item:selected { background: $sel_list; color: $text; }
QTableWidget { background: $card; border: 1px solid $border; border-radius: 10px; gridline-color: $grid;
               selection-background-color: $sel; selection-color: $text; }
QHeaderView::section { background: $header; border: none; border-bottom: 1px solid $border;
                       border-right: 1px solid $border; padding: 4px 6px; font-weight: 600; }
QTabWidget::pane { border: 1px solid $border; border-radius: 8px; background: $card; }
QTabBar::tab { background: $tab; padding: 7px 14px; border-top-left-radius: 8px; border-top-right-radius: 8px; margin-right: 2px; }
QTabBar::tab:selected { background: $card; font-weight: 600; }
QSplitter::handle { background: transparent; }
""")


def apply_theme(app, name: str):
    global THEME_NAME, THEME
    THEME_NAME, THEME = name, THEMES[name]
    t = THEME
    pal = QPalette()
    R = QPalette.ColorRole
    for role, col in ((R.Window, t["bg"]), (R.WindowText, t["text"]), (R.Base, t["card"]), (R.AlternateBase, t["bg"]),
                      (R.Text, t["text"]), (R.Button, t["card"]), (R.ButtonText, t["text"]),
                      (R.ToolTipBase, t["card"]), (R.ToolTipText, t["text"]), (R.Highlight, t["sel"]),
                      (R.HighlightedText, t["text"]), (R.PlaceholderText, t["muted"])):
        pal.setColor(role, QColor(col))
    app.setPalette(pal)
    app.setStyleSheet(QSS.substitute(t))


def default_theme() -> str:
    saved = load_settings().get("theme")
    if saved in THEMES:
        return saved
    try:
        return "dark" if QGuiApplication.styleHints().colorScheme() == Qt.ColorScheme.Dark else "light"
    except Exception:
        return "light"


# ════════════════════════════════════════════════════════════════════════════
#  5) GUI
# ════════════════════════════════════════════════════════════════════════════
from PySide6.QtCore import Qt, QUrl, Signal
from PySide6.QtGui import QBrush, QColor, QDesktopServices, QGuiApplication, QIcon, QKeySequence, QPalette, QShortcut
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFileDialog, QFrame,
    QGridLayout, QHBoxLayout, QHeaderView, QInputDialog, QLabel, QLineEdit, QListWidget, QListWidgetItem,
    QMenu, QMessageBox, QPushButton, QSizePolicy, QSpinBox, QSplitter, QStackedWidget, QTableWidget, QTableWidgetItem,
    QTabWidget, QVBoxLayout, QWidget,
)

COL_CHK, COL_N, COL_SRC, COL_DST, COL_MODE, COL_VAL, COL_STAT = range(7)
HEADERS = ["✓", "#", "Source cell(s)", "Target cell", "Mode", "Value to write", "Status"]
HINT_NORMAL = ("Select cells (click / drag / Ctrl+click), then send them to Source or Target →   "
               "Blue = source · Green = target")


def card() -> QFrame:
    f = QFrame()
    f.setObjectName("card")
    return f


def _ro(text: str, center: bool = False) -> QTableWidgetItem:
    it = QTableWidgetItem(text)
    it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
    if center:
        it.setTextAlignment(Qt.AlignCenter)
    return it


def _h(text: str) -> QLabel:
    l = QLabel(text)
    l.setObjectName("h")
    return l


class SheetMapDialog(QDialog):
    """Asks, for every sheet the mappings use, which sheet of the current file should be used instead."""

    def __init__(self, tpl_name: str, usage: dict, keys: list[str], main_count: int, suggestions: dict,
                 file_label: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Match sheets")
        self.setMinimumWidth(580)
        lay = QVBoxLayout(self)
        info = QLabel(f"«{tpl_name}» uses sheets that don't all exist in «{file_label}».\n"
                      "For each one, choose which sheet to use here (a suggestion is pre-selected):")
        info.setWordWrap(True)
        lay.addWidget(info)
        g = QGridLayout()
        g.addWidget(_h("Sheet in the mappings"), 0, 0)
        g.addWidget(_h("Use this sheet instead"), 0, 1)
        self.combos: dict[str, QComboBox] = {}
        for i, (sh, (ns, nd)) in enumerate(usage.items(), start=1):
            lbl = QLabel(f"{key_label(sh)}\n{ns} source · {nd} target cell(s)")
            lbl.setObjectName("h" if sh in keys else "muted")
            cb = QComboBox()
            cb.addItem("— skip (leave these mappings unmatched) —", "")
            for k in (keys[:main_count] if nd else keys):       # targets must be sheets of this file
                cb.addItem(key_label(k), k)
            cb.setCurrentIndex(max(cb.findData(suggestions.get(sh, "")), 0))
            self.combos[sh] = cb
            g.addWidget(lbl, i, 0)
            g.addWidget(cb, i, 1)
        g.setColumnStretch(1, 1)
        lay.addLayout(g)
        if any(sh.startswith("[") for sh in usage):
            tip = QLabel("Tip: to read from another file, press Cancel and attach it first with «Source files».")
            tip.setObjectName("muted")
            tip.setWordWrap(True)
            lay.addWidget(tip)
        bb = QDialogButtonBox()
        ok = bb.addButton("Apply", QDialogButtonBox.AcceptRole)
        ok.setObjectName("primary")
        bb.addButton("Cancel", QDialogButtonBox.RejectRole)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def result_map(self) -> dict[str, str]:
        return {k: (cb.currentData() or "") for k, cb in self.combos.items()}


def remap_rules(parent, tpl_name: str, book: Book, rules: list[dict], index: Optional[dict] = None,
                alias: Optional[dict] = None, file_label: str = "") -> Optional[tuple[list[dict], dict, dict]]:
    """Makes the sheet names in `rules` fit `book`. Silent when every sheet already exists; otherwise asks.
    Returns (new_rules, choice, smap) or None if the user cancelled."""
    tpl_sheets = list(dict.fromkeys([s[0] for r in rules for s in r["src"]] + [r["dst"][0] for r in rules]))
    smap: dict[str, str] = {}
    choice: dict[str, str] = {}
    if any(not book.has(sh) for sh in tpl_sheets):
        index, alias = index or {}, alias or {}
        keys = book.all_keys()
        usage = {sh: (sum(s[0] == sh for r in rules for s in r["src"]), sum(r["dst"][0] == sh for r in rules))
                 for sh in tpl_sheets}
        sugg = {sh: suggest_sheet(sh, keys, index.get(sh), alias.get(sh), len(book.main_sheets())) for sh in tpl_sheets}
        dlg = SheetMapDialog(tpl_name, usage, keys, len(book.main_sheets()), sugg, file_label or book.path.name, parent)
        if dlg.exec() != QDialog.Accepted:
            return None
        choice = dlg.result_map()
        smap = {k: v for k, v in choice.items() if v and v != k}
    new = deepcopy(rules)
    for r in new:
        r["src"] = [[smap.get(s[0], s[0]), s[1]] for s in r["src"]]
        r["dst"] = [smap.get(r["dst"][0], r["dst"][0]), r["dst"][1]]
    return new, choice, smap


class PairPreviewDialog(QDialog):
    """Shows exactly which source cell goes to which target cell before the mapping is added."""

    def __init__(self, book: Book, rules: list[dict], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Check the pairing")
        self.resize(780, 480)
        lay = QVBoxLayout(self)
        info = QLabel("Each source cell is copied to the target cell on the same row. Check that the pairs are right:")
        info.setWordWrap(True)
        lay.addWidget(info)
        t = QTableWidget(len(rules), 4)
        t.setHorizontalHeaderLabels(["#", "Source", "Value", "→ Target"])
        t.verticalHeader().setVisible(False)
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        t.setSelectionBehavior(QAbstractItemView.SelectRows)
        for i, r in enumerate(rules):
            val = fmt_display(rule_value(book, r, None, True)).replace("\n", " ⏎ ")
            t.setItem(i, 0, _ro(str(i + 1), True))
            t.setItem(i, 1, _ro("; ".join(fmt_ref(*s) for s in r["src"])))
            t.setItem(i, 2, _ro(val))
            t.setItem(i, 3, _ro(fmt_ref(*r["dst"])))
        hh = t.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.Stretch)
        hh.setSectionResizeMode(0, QHeaderView.Fixed)
        t.setColumnWidth(0, 44)
        lay.addWidget(t, 1)
        bb = QDialogButtonBox()
        ok = bb.addButton("Add mapping", QDialogButtonBox.AcceptRole)
        ok.setObjectName("primary")
        bb.addButton("Cancel", QDialogButtonBox.RejectRole)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)


class FilePickDialog(QDialog):
    def __init__(self, panels: list, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Choose files")
        self.setMinimumWidth(480)
        lay = QVBoxLayout(self)
        lab = QLabel("Apply the same operation to these files:")
        lab.setObjectName("h")
        lay.addWidget(lab)
        self.lst = QListWidget()
        for p in panels:
            it = QListWidgetItem(p.path.name)
            it.setToolTip(str(p.path))
            it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked)
            it.setData(Qt.UserRole, p)
            self.lst.addItem(it)
        lay.addWidget(self.lst, 1)
        bb = QDialogButtonBox()
        ok = bb.addButton("Continue", QDialogButtonBox.AcceptRole)
        ok.setObjectName("primary")
        bb.addButton("Cancel", QDialogButtonBox.RejectRole)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def chosen(self) -> list:
        return [self.lst.item(i).data(Qt.UserRole) for i in range(self.lst.count())
                if self.lst.item(i).checkState() == Qt.Checked]


class SeqSheetDialog(QDialog):
    """Lets the user choose, for ONE record file, which of its sheets holds the data."""

    def __init__(self, file_name: str, used: list[str], ref_sheets: list[str], rec_sheets: list[str],
                 current: dict, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Choose sheets")
        self.setMinimumWidth(500)
        lay = QVBoxLayout(self)
        info = QLabel(f"Which sheet of «{file_name}» holds the data for each sheet of the sample file?")
        info.setWordWrap(True)
        lay.addWidget(info)
        g = QGridLayout()
        g.addWidget(_h("Sheet in the sample file"), 0, 0)
        g.addWidget(_h(f"Sheet to use in {file_name}"), 0, 1)
        self.combos: dict[str, QComboBox] = {}
        for i, sh in enumerate(used, start=1):
            g.addWidget(QLabel(sh), i, 0)
            cb = QComboBox()
            cb.addItems(rec_sheets)
            cb.setCurrentIndex(max(cb.findText(_resolve_one(sh, ref_sheets, rec_sheets, current.get(sh))), 0))
            self.combos[sh] = cb
            g.addWidget(cb, i, 1)
        g.setColumnStretch(1, 1)
        lay.addLayout(g)
        self.chk_all = QCheckBox("Use the same choice for the other files that have a sheet with this name")
        self.chk_all.setChecked(True)
        lay.addWidget(self.chk_all)
        bb = QDialogButtonBox()
        ok = bb.addButton("Apply", QDialogButtonBox.AcceptRole)
        ok.setObjectName("primary")
        bb.addButton("Cancel", QDialogButtonBox.RejectRole)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def result_map(self) -> dict[str, str]:
        return {k: cb.currentText() for k, cb in self.combos.items()}

    def apply_all(self) -> bool:
        return self.chk_all.isChecked()


class SequentialDialog(QDialog):
    """Fills the target table record by record from many source files (one file = one row / column)."""

    def __init__(self, panel: "FilePanel", labels: list[str]):
        super().__init__(panel)
        self.p = panel
        self.plan: Optional[dict] = None
        self.overrides: dict[str, dict[str, str]] = {}      # record file -> {sample sheet: sheet to use}
        seq = panel.seq
        self.setWindowTitle("Sequential fill")
        scr = QApplication.primaryScreen().availableGeometry()
        self.resize(min(1020, int(scr.width() * 0.92)), min(720, int(scr.height() * 0.92)))
        self.setMinimumSize(min(720, scr.width()), min(520, scr.height()))
        lay = QVBoxLayout(self)
        info = QLabel(f"Fills «{panel.path.name}» record by record. The cells you mapped are read from EACH source file "
                      "in the list, and every file goes to the next row (or column) of the target.")
        info.setWordWrap(True)
        lay.addWidget(info)
        top = QHBoxLayout()
        lay.addLayout(top)

        fc = card()
        fl = QVBoxLayout(fc)
        fl.setContentsMargins(10, 8, 10, 8)
        fl.addWidget(_h("Source files, in filling order (drag to reorder)"))
        self.lst = QListWidget()
        self.lst.setDragDropMode(QAbstractItemView.InternalMove)
        self.lst.setMinimumHeight(110)
        self.lst.setMaximumHeight(190)
        self.lst.model().rowsMoved.connect(lambda *_: self._invalidate())
        fl.addWidget(self.lst)
        br = QHBoxLayout()
        for txt, fn in (("Add files…", self.add_files), ("Add folder…", self.add_folder)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            br.addWidget(b)
        self.btn_open = QPushButton("From opened files ▾")
        self.menu_open = QMenu(self)
        self.menu_open.aboutToShow.connect(self._build_open_menu)
        self.btn_open.setMenu(self.menu_open)
        br.addWidget(self.btn_open)
        fl.addLayout(br)
        br2 = QHBoxLayout()
        for txt, fn in (("Remove", self.remove_selected), ("Sort A→Z", self.sort_az), ("▲", lambda: self.move(-1)),
                        ("▼", lambda: self.move(1)), ("Clear", self.clear_files)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            br2.addWidget(b)
        self.lbl_count = QLabel("")
        self.lbl_count.setObjectName("muted")
        br2.addWidget(self.lbl_count, 1)
        fl.addLayout(br2)
        top.addWidget(fc, 3)

        oc = card()
        g = QGridLayout(oc)
        g.setContentsMargins(10, 8, 10, 8)
        g.addWidget(_h("Sample file:"), 0, 0)
        self.cmb_ref = QComboBox()
        self.cmb_ref.addItems(labels)
        self.cmb_ref.setToolTip("The source file your mappings were made on")
        g.addWidget(self.cmb_ref, 0, 1)
        g.addWidget(_h("Next record goes:"), 1, 0)
        self.cmb_dir = QComboBox()
        self.cmb_dir.addItem("Next row  (↓)", "down")
        self.cmb_dir.addItem("Next column  (→)", "right")
        self.cmb_dir.setCurrentIndex(max(self.cmb_dir.findData(seq.get("direction", "down")), 0))
        g.addWidget(self.cmb_dir, 1, 1)
        g.addWidget(_h("Step:"), 2, 0)
        self.spn_step = QSpinBox()
        self.spn_step.setRange(1, 100)
        self.spn_step.setValue(int(seq.get("step", 1)))
        self.spn_step.setToolTip("Rows (or columns) between two records. Use 1 for a normal table.")
        g.addWidget(self.spn_step, 2, 1)
        self.chk_auto = QCheckBox("Start at the first empty row/column")
        self.chk_auto.setChecked(bool(seq.get("auto", True)))
        self.chk_auto.setToolTip("Skips records already in the table, so you can add new ones below them")
        g.addWidget(self.chk_auto, 3, 0, 1, 2)
        self.chk_counter = QCheckBox("Number the records in cell:")
        self.chk_counter.setChecked(bool(seq.get("counter_on", False)))
        g.addWidget(self.chk_counter, 4, 0)
        self.ed_counter = QLineEdit(seq.get("counter_cell", self._guess_counter()))
        self.ed_counter.setPlaceholderText("e.g. Sheet1!A2 (first record)")
        g.addWidget(self.ed_counter, 4, 1)
        g.addWidget(_h("Start number:"), 5, 0)
        self.spn_start = QSpinBox()
        self.spn_start.setRange(0, 1000000)
        self.spn_start.setSpecialValueText("Auto")
        self.spn_start.setValue(int(seq.get("counter_start", 0)))
        self.spn_start.setToolTip("Auto = continues after the records already in the table")
        g.addWidget(self.spn_start, 5, 1)
        g.setColumnStretch(1, 1)
        top.addWidget(oc, 2)

        self.lbl_info = QLabel("Add the source files, then press Preview to check the result.")
        self.lbl_info.setObjectName("muted")
        self.lbl_info.setWordWrap(True)
        lay.addWidget(self.lbl_info)
        self.tbl = QTableWidget(0, 0)
        self.tbl.verticalHeader().setVisible(False)
        self.tbl.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.tbl.cellDoubleClicked.connect(self._edit_sheets)
        self.tbl.setToolTip("Double-click a file to choose which of its sheets is used")
        lay.addWidget(self.tbl, 1)

        bb = QHBoxLayout()
        self.btn_prev = QPushButton("Preview")
        self.btn_prev.clicked.connect(self.preview)
        self.btn_exec = QPushButton("Execute && Save new copy")
        self.btn_exec.setObjectName("success")
        self.btn_exec.setEnabled(False)
        self.btn_exec.clicked.connect(self.execute)
        b_cancel = QPushButton("Close")
        b_cancel.clicked.connect(self.reject)
        bb.addWidget(self.btn_prev)
        bb.addStretch(1)
        bb.addWidget(b_cancel)
        bb.addWidget(self.btn_exec)
        lay.addLayout(bb)

        ref_path = panel.book.ext_paths.get(self.cmb_ref.currentText())
        if ref_path:                                 # the sample is usually one of the records too
            self._add_paths([ref_path])
        for w_ in (self.cmb_ref, self.cmb_dir):
            w_.currentIndexChanged.connect(self._invalidate)
        for w_ in (self.spn_step, self.spn_start):
            w_.valueChanged.connect(self._invalidate)
        for w_ in (self.chk_auto, self.chk_counter):
            w_.toggled.connect(self._invalidate)
        self.ed_counter.textChanged.connect(self._invalidate)
        self._refresh_count()

    def _guess_counter(self) -> str:
        for r in self.p.rules:
            if r["enabled"] and any(s[0].startswith("[") for s in r["src"]):
                return fmt_ref(r["dst"][0], f"A{split_coord(r['dst'][1])[0]}")
        return ""

    # files list
    def _files(self) -> list[Path]:
        return [Path(self.lst.item(i).data(Qt.UserRole)) for i in range(self.lst.count())]

    def _set_files(self, files: list[Path]):
        self.lst.clear()
        for f in files:
            it = QListWidgetItem(f.name)
            it.setToolTip(str(f))
            it.setData(Qt.UserRole, str(f))
            self.lst.addItem(it)
        self._invalidate()
        self._refresh_count()

    def _add_paths(self, paths):
        have = {str(f.resolve()) for f in self._files()}
        new = list(self._files())
        for p in paths:
            p = Path(p)
            if p.suffix.lower() not in (".xlsx", ".xlsm") or p.name.startswith("~$"):
                continue
            if p.resolve() == self.p.path.resolve() or str(p.resolve()) in have:
                continue
            have.add(str(p.resolve()))
            new.append(p)
        self._set_files(new)

    def _refresh_count(self):
        self.lbl_count.setText(f"{self.lst.count()} file(s)")

    def add_files(self):
        fns, _ = QFileDialog.getOpenFileNames(self, "Choose source files", str(self.p.path.parent), "Excel (*.xlsx *.xlsm)")
        self._add_paths(fns)

    def add_folder(self):
        d = QFileDialog.getExistingDirectory(self, "Choose a folder with the source files", str(self.p.path.parent))
        if d:
            self._add_paths(sorted(Path(d).glob("*.xls[xm]"), key=lambda x: x.name.lower()))

    def _build_open_menu(self):
        m = self.menu_open
        m.clear()
        have = {str(f.resolve()) for f in self._files()}
        others = [o for o in self.p.others_provider() if str(o.path.resolve()) not in have]
        for o in others:
            m.addAction(o.path.name).triggered.connect(lambda _=False, pp=o.path: self._add_paths([pp]))
        if len(others) > 1:
            m.addSeparator()
            m.addAction("Add all").triggered.connect(lambda _=False: self._add_paths([o.path for o in others]))
        if not others:
            m.addAction("(no other open files)").setEnabled(False)

    def remove_selected(self):
        for it in self.lst.selectedItems():
            self.lst.takeItem(self.lst.row(it))
        self._invalidate()
        self._refresh_count()

    def clear_files(self):
        self._set_files([])

    def sort_az(self):
        self._set_files(sorted(self._files(), key=lambda f: f.name.lower()))

    def move(self, d: int):
        i = self.lst.currentRow()
        j = i + d
        if i < 0 or not 0 <= j < self.lst.count():
            return
        files = self._files()
        files[i], files[j] = files[j], files[i]
        self._set_files(files)
        self.lst.setCurrentRow(j)

    # plan / preview
    def _invalidate(self, *_):
        self.plan = None
        self.btn_exec.setEnabled(False)

    def _counter(self) -> Optional[tuple[str, str, int]]:
        if not self.chk_counter.isChecked():
            return None
        cells = parse_refs(self.ed_counter.text(), self.p.rules[0]["dst"][0])
        if len(cells) != 1 or not self.p.book.is_main(cells[0][0]):
            raise ValueError("The record number must go in exactly one cell of this file's own sheets (the first record's cell).")
        return cells[0][0], cells[0][1], self.spn_start.value()

    def _store(self):
        new = {"direction": self.cmb_dir.currentData(), "step": self.spn_step.value(), "auto": self.chk_auto.isChecked(),
               "counter_on": self.chk_counter.isChecked(), "counter_cell": self.ed_counter.text().strip(),
               "counter_start": self.spn_start.value()}
        if new != self.p.seq:
            self.p.seq = new
            if self.p.cur_template:
                self.p.dirty = True
                self.p._update_tpl_label()

    def done(self, r):
        self._store()
        super().done(r)

    def preview(self):
        files = self._files()
        if not files:
            QMessageBox.information(self, APP_NAME, "Add the source files first.")
            return
        try:
            counter = self._counter()
        except ValueError as e:
            QMessageBox.warning(self, APP_NAME, str(e))
            return
        self._store()

        def prog(n, total, name):
            self.lbl_info.setText(f"Reading {n}/{total}: {name}")
            QApplication.processEvents()

        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            plan = build_sequence(self.p.book, self.p.rules, self.cmb_ref.currentText(), files, self.cmb_dir.currentData(),
                                  self.spn_step.value(), self.chk_auto.isChecked(), counter, self.p.chk_over.isChecked(),
                                  self.p.chk_digits.isChecked(), prog, self.overrides)
        except Exception as ex:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, APP_NAME, f"Could not build the preview:\n\n{traceback.format_exc()[-900:]}")
            return
        QApplication.restoreOverrideCursor()
        self.plan = plan
        self._fill_preview(plan)
        skipped = len(plan["records"]) - plan["n_ok"]
        where = "row" if plan["unit"][0] else "column"
        self.lbl_info.setText(f"{plan['n_ok']} record(s) ready · {skipped} skipped"
                              + (f" · {len(plan['warnings'])} with ⚠ warnings (check them)" if plan["warnings"] else "")
                              + f" · {len(plan['writes'])} cells to write · double-click a file to change the sheet it uses"
                              + (f" · starting {plan['offset'] * max(plan['unit'])} {where}(s) after the first position"
                                 if plan["offset"] else ""))
        self.btn_exec.setEnabled(plan["n_ok"] > 0)

    def _fill_preview(self, plan: dict):
        cols = ["File"] + (["#"] if plan["counter"] else []) + [f"{co}  {hd}".strip() for _, co, hd in plan["columns"]]
        t = self.tbl
        t.clear()
        t.setColumnCount(len(cols))
        t.setRowCount(len(plan["records"]))
        t.setHorizontalHeaderLabels(cols)
        u = plan["unit"]
        for row, (path, status, writes, k) in enumerate(plan["records"]):
            if status != "ok":
                it = _ro(f"✗ {path.name} — {status}")
                it.setBackground(QBrush(QColor(THEME["bad"])))
                t.setItem(row, 0, it)
                continue
            vals = {(w[0], w[1]): w[2] for w in writes}
            warn = plan["warnings"].get(str(path))
            f_it = _ro(f"⚠ {path.name} — {warn}" if warn else path.name)
            if warn:
                f_it.setBackground(QBrush(QColor(THEME["warn"])))
            t.setItem(row, 0, f_it)
            c = 1
            if plan["counter"]:
                cs, cc, _ = plan["counter"]
                t.setItem(row, c, _ro(fmt_display(vals.get((cs, shift_coord(cc, u[0] * k, u[1] * k))))))
                c += 1
            for sh, co, _ in plan["columns"]:
                t.setItem(row, c, _ro(fmt_display(vals.get((sh, shift_coord(co, u[0] * k, u[1] * k)))).replace("\n", " ⏎ ")))
                c += 1
        t.resizeColumnsToContents()
        t.horizontalHeader().setStretchLastSection(True)

    def _edit_sheets(self, row: int, _col: int):
        if not self.plan or row >= len(self.plan["records"]):
            return
        path = self.plan["records"][row][0]
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            rec_sheets = list(wb.sheetnames)
            wb.close()
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, f"Could not open {path.name}:\n{ex}")
            return
        dlg = SeqSheetDialog(path.name, self.plan["used_sheets"], self.plan["ref_sheets"], rec_sheets,
                             self.overrides.get(str(path), {}), self)
        if dlg.exec() != QDialog.Accepted:
            return
        choice = dlg.result_map()
        self.overrides[str(path)] = choice
        if dlg.apply_all():                      # other files that would otherwise have to guess the same sheet
            for f in self._files():
                if f == path:
                    continue
                try:
                    w2 = openpyxl.load_workbook(f, read_only=True)
                    names = list(w2.sheetnames)
                    w2.close()
                except Exception:
                    continue
                for sh, rs in choice.items():
                    if rs in names and sh not in names:
                        self.overrides.setdefault(str(f), {})[sh] = rs
        self.preview()

    def execute(self):
        plan = self.plan
        if not plan or not plan["n_ok"]:
            return
        if QMessageBox.question(self, APP_NAME, f"Fill {plan['n_ok']} record(s) ({len(plan['writes'])} cells) into "
                                                f"«{self.p.path.name}»?\nA new copy is saved next to the original.") != QMessageBox.Yes:
            return
        try:
            out, n = self.p._write_file(self.p.path, plan["writes"], unique_output_path(self.p.path))
        except _Cancelled:
            return
        except Exception:
            QMessageBox.critical(self, APP_NAME, f"Save failed:\n\n{traceback.format_exc()[-1200:]}")
            return
        skipped = [f"✗ {r[0].name} — {r[1]}" for r in plan["records"] if r[1] != "ok"]
        box = QMessageBox(self)
        box.setWindowTitle("Done")
        box.setIcon(QMessageBox.Information)
        box.setText(f"<b>{plan['n_ok']}</b> record(s) filled ({n} cells). File saved:\n{out}")
        warned = [f"⚠ {Path(k).name} — {v}" for k, v in plan["warnings"].items()]
        if skipped or warned:
            box.setInformativeText("\n".join(([("Skipped:")] + skipped if skipped else []) + warned))
        op = box.addButton("Open folder", QMessageBox.ActionRole)
        box.addButton("Close", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is op:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(out.parent)))
        self.accept()


class FilePanel(QWidget):
    """Fully independent panel per Excel file."""
    templates_changed = Signal()

    def __init__(self, path: Path, templates: Templates, parent=None):
        super().__init__(parent)
        self.path = path
        self.templates = templates
        self.book = Book(path)
        self.others_provider: Callable[[], list] = lambda: []      # set by the main window
        self.rules: list[dict] = []
        self.overrides: dict[str, str] = {}      # manual values for THIS file only (rule id -> text)
        self.cur_template = ""
        self._tpl_sheet_map: dict[str, str] = {}   # template sheet name -> sheet of THIS file
        self.dirty = False
        self.seq: dict = {}                       # sequential-fill settings (saved with the template)
        self.views: dict[str, QTableWidget] = {}
        self._painted: set[tuple[str, int, int]] = set()
        self._loading = False
        self._pair_src: Optional[tuple[str, str]] = None      # pending source in «pair by clicking» mode
        self._build_ui()
        esc = QShortcut(QKeySequence("Esc"), self)
        esc.activated.connect(self._pair_reset)
        self.refresh_templates()
        self.refresh_rules()

    # ── UI ──────────────────────────────────────────────────────────────────
    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(10)

        top = card()
        tl = QVBoxLayout(top)
        tl.setContentsMargins(14, 8, 14, 8)
        nrow = QHBoxLayout()
        name = QLabel(self.path.name)
        name.setObjectName("title")
        loc = QLabel(str(self.path.parent))
        loc.setObjectName("muted")
        loc.setToolTip(str(self.path.parent))
        loc.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        for w_ in (name, loc):                     # long names/paths must not force a wide window
            w_.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        nrow.addWidget(name, 2)
        nrow.addWidget(loc, 3)
        tl.addLayout(nrow)

        r1 = QHBoxLayout()
        r1.addWidget(_h("Template:"))
        self.cmb_tpl = QComboBox()
        self.cmb_tpl.setMinimumWidth(240)
        self.cmb_tpl.currentIndexChanged.connect(self._tpl_chosen)
        r1.addWidget(self.cmb_tpl)
        b_save = QPushButton("Save")
        b_save.clicked.connect(self.save_template)
        b_as = QPushButton("Save as new…")
        b_as.clicked.connect(self.save_template_as)
        b_del = QPushButton("Delete template")
        b_del.clicked.connect(self.delete_template)
        for b in (b_save, b_as, b_del):
            r1.addWidget(b)
        self.lbl_tpl = QLabel("")
        self.lbl_tpl.setObjectName("muted")
        r1.addWidget(self.lbl_tpl, 1)
        tl.addLayout(r1)

        r2 = QHBoxLayout()
        self.chk_over = QCheckBox("Overwrite non-empty target cells")
        self.chk_over.setChecked(True)
        self.chk_digits = QCheckBox("Convert Arabic-Indic digits to Latin (350)")
        self.chk_digits.setChecked(True)
        for c in (self.chk_over, self.chk_digits):
            c.toggled.connect(lambda _=None: self.refresh_rules())
            r2.addWidget(c)
        r2.addStretch(1)
        tl.addLayout(r2)
        root.addWidget(top)

        split = QSplitter(Qt.Horizontal)
        root.addWidget(split, 1)

        # left: sheet preview (this file + attached source files)
        left = QWidget()
        ll = QVBoxLayout(left)
        ll.setContentsMargins(0, 0, 0, 0)
        hr = QHBoxLayout()
        self.lbl_hint = QLabel(HINT_NORMAL)
        self.lbl_hint.setObjectName("muted")
        self.lbl_hint.setWordWrap(True)
        hr.addWidget(self.lbl_hint, 1)
        self.btn_pair = QPushButton("🔗 Pair by clicking")
        self.btn_pair.setCheckable(True)
        self.btn_pair.setToolTip("Click a source cell, then the target cell it must go to. "
                                 "Each pair is added immediately, in exactly the order you choose.")
        self.btn_pair.toggled.connect(self._pair_toggled)
        hr.addWidget(self.btn_pair)
        self.btn_srcfiles = QPushButton("Source files ▾")
        self.btn_srcfiles.setToolTip("Use cells from other Excel files as sources")
        self.src_menu = QMenu(self)
        self.src_menu.aboutToShow.connect(self._build_src_menu)
        self.btn_srcfiles.setMenu(self.src_menu)
        hr.addWidget(self.btn_srcfiles)
        ll.addLayout(hr)
        self.tabs = QTabWidget()
        for n in self.book.main_sheets():
            self._add_view(n)
        ll.addWidget(self.tabs, 1)
        split.addWidget(left)

        # right: builder + mappings
        right = QWidget()
        rl = QVBoxLayout(right)
        rl.setContentsMargins(0, 0, 0, 0)
        rl.setSpacing(8)

        b = card()
        g = QGridLayout(b)
        g.setContentsMargins(12, 10, 12, 10)
        self.ed_src, self.ed_dst = QLineEdit(), QLineEdit()
        self.ed_src.setPlaceholderText("e.g.  Survey!B5:B40 ; '[Other.xlsx]Info'!C3")
        self.ed_dst.setPlaceholderText("e.g.  Report!C10:C45")
        for row, (txt, ed) in enumerate((("Source:", self.ed_src), ("Target:", self.ed_dst))):
            g.addWidget(_h(txt), row, 0)
            g.addWidget(ed, row, 1)
            bs = QPushButton("◂ Set")
            bs.setToolTip("Replace with the current selection in the sheet preview")
            bs.clicked.connect(lambda _=False, e=ed: self.take_selection(e, False))
            ba = QPushButton("＋ Add")
            ba.setToolTip("Append the current selection (useful for cells from another sheet or file)")
            ba.clicked.connect(lambda _=False, e=ed: self.take_selection(e, True))
            g.addWidget(bs, row, 2)
            g.addWidget(ba, row, 3)
        g.addWidget(_h("Pairing:"), 2, 0)
        h2 = QHBoxLayout()
        self.cmb_order = QComboBox()
        self.cmb_order.addItem("Row by row (→ then ↓)", "row")
        self.cmb_order.addItem("Column by column (↓ then →)", "col")
        self.cmb_order.setToolTip("Order in which the cells of a range are paired (same order on both sides)")
        self.cmb_mode = QComboBox()
        for k, v in MODES.items():
            self.cmb_mode.addItem(f"Many → one cell: {v}", k)
        self.cmb_mode.setToolTip("Only used when several source cells feed a single target cell")
        h2.addWidget(self.cmb_order, 1)
        h2.addWidget(self.cmb_mode, 1)
        g.addLayout(h2, 2, 1, 1, 3)
        self.chk_preview = QCheckBox("Preview the pairs before adding")
        self.chk_preview.setChecked(True)
        g.addWidget(self.chk_preview, 3, 0, 1, 2)
        b_add = QPushButton("Add mapping")
        b_add.setObjectName("primary")
        b_add.clicked.connect(self.add_mapping)
        g.addWidget(b_add, 3, 2, 1, 2)
        g.setColumnStretch(1, 1)
        rl.addWidget(b)

        self.tbl = QTableWidget(0, len(HEADERS))
        self.tbl.setHorizontalHeaderLabels(HEADERS)
        self.tbl.verticalHeader().setVisible(False)
        self.tbl.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tbl.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed)
        self.tbl.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tbl.customContextMenuRequested.connect(self._ctx_menu)
        self.tbl.cellDoubleClicked.connect(self._dbl)
        self.tbl.itemChanged.connect(self._item_changed)
        self.tbl.setToolTip("Double-click Source/Target to edit the cells · Mode to change it · "
                            "Value to override it for this file only · right-click for more")
        hh = self.tbl.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.Interactive)
        for c in (COL_SRC, COL_DST, COL_VAL):
            hh.setSectionResizeMode(c, QHeaderView.Stretch)
        for c, w in ((COL_CHK, 40), (COL_N, 40), (COL_MODE, 90), (COL_STAT, 110)):
            self.tbl.setColumnWidth(c, w)
        rl.addWidget(self.tbl, 1)

        bar = QHBoxLayout()
        for txt, fn in (("▲", lambda: self.move(-1)), ("▼", lambda: self.move(1)),
                        ("Delete selected", self.delete_selected), ("Clear all", self.clear_all)):
            bt = QPushButton(txt)
            bt.clicked.connect(fn)
            bar.addWidget(bt)
        bar.addStretch(1)
        self.lbl_sum = QLabel("")
        self.lbl_sum.setObjectName("muted")
        bar.addWidget(self.lbl_sum)
        rl.addLayout(bar)

        self.btn_run = QPushButton("Execute && Save new copy")
        self.btn_run.setObjectName("success")
        self.btn_run.clicked.connect(self.run_export)
        self.btn_seq = QPushButton("Sequential fill…")
        self.btn_seq.setObjectName("primary")
        self.btn_seq.setToolTip("Fill a table record by record: the same cells are read from many source files, one file per row")
        self.btn_seq.clicked.connect(self.open_sequential)
        brow = QHBoxLayout()
        brow.addWidget(self.btn_seq)
        brow.addWidget(self.btn_run, 1)
        rl.addLayout(brow)

        split.addWidget(right)
        split.setStretchFactor(0, 3)
        split.setStretchFactor(1, 2)
        split.setSizes([760, 600])

    def on_theme_changed(self):
        self.refresh_rules()                     # re-applies status colours and sheet highlights

    # ── sheet views ─────────────────────────────────────────────────────────
    def _add_view(self, key: str):
        t = self._build_view(key)
        t.setProperty("key", key)
        self.views[key] = t
        t.cellClicked.connect(lambda r, c, tt=t: self._cell_clicked(tt, r, c))
        i = self.tabs.addTab(t, key_label(key))
        self.tabs.setTabToolTip(i, key_label(key))

    def _build_view(self, key: str) -> QTableWidget:
        ws = self.book.ws(key)
        nr = min(max(ws.max_row or 1, 1), MAX_ROWS)
        nc = min(max(ws.max_column or 1, 1), MAX_COLS)
        t = QTableWidget(nr + 20, nc + 3)          # a few extra empty rows/cols so empty targets can be selected
        t.setHorizontalHeaderLabels([get_column_letter(i + 1) for i in range(nc + 3)])
        t.horizontalHeader().setDefaultSectionSize(120)
        t.setSelectionMode(QAbstractItemView.ExtendedSelection)
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        t.setWordWrap(False)
        for row in ws.iter_rows(min_row=1, max_row=nr, max_col=nc):
            for c in row:
                if is_blank(c.value):
                    continue
                txt = fmt_display(c.value)
                it = QTableWidgetItem(txt.replace("\n", " ⏎ "))
                it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
                if len(txt) > 30:
                    it.setToolTip(txt)
                t.setItem(c.row - 1, c.column - 1, it)
        for rng in ws.merged_cells.ranges:
            if rng.max_row <= nr and rng.max_col <= nc and (rng.max_row > rng.min_row or rng.max_col > rng.min_col):
                t.setSpan(rng.min_row - 1, rng.min_col - 1, rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        return t

    def _cur_key(self) -> str:
        w = self.tabs.currentWidget()
        return w.property("key") if w is not None else ""

    # ── source files (other workbooks) ──────────────────────────────────────
    def _build_src_menu(self):
        m = self.src_menu
        m.clear()
        attached = {str(Path(p).resolve()) for p in self.book.ext_paths.values()}
        others = [p for p in self.others_provider() if str(p.path.resolve()) not in attached]
        sub = m.addMenu("Add from opened files")
        if others:
            for p in others:
                sub.addAction(p.path.name).triggered.connect(lambda _=False, pp=p.path: self.add_source_file(pp))
        else:
            sub.setEnabled(False)
        m.addAction("Browse for a file…").triggered.connect(self.browse_source)
        if self.book.ext:
            m.addSeparator()
            for lab in self.book.ext:
                m.addAction(f"Remove «{lab}»").triggered.connect(lambda _=False, l=lab: self.remove_source_file(l))

    def browse_source(self):
        fns, _ = QFileDialog.getOpenFileNames(self, "Choose source Excel files", str(self.path.parent),
                                              "Excel (*.xlsx *.xlsm)")
        for fn in fns:
            self.add_source_file(Path(fn))

    def add_source_file(self, path: Path):
        path = Path(path)
        if path.resolve() == self.path.resolve():
            QMessageBox.information(self, APP_NAME, "This is the file itself; its sheets are already available as sources.")
            return
        if any(Path(p).resolve() == path.resolve() for p in self.book.ext_paths.values()):
            return
        try:
            lab = self.book.add_external(path)
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, f"Could not open {path.name}:\n{ex}")
            return
        for n in self.book.ext[lab].sheetnames:
            self._add_view(f"[{lab}]{n}")
        self._update_srcbtn()
        self.refresh_rules()

    def remove_source_file(self, lab: str):
        used = sum(1 for r in self.rules for s in r["src"] if self.book.split_key(s[0])[0] == lab)
        if used and QMessageBox.question(
                self, APP_NAME, f"{used} source cell(s) in your mappings come from «{lab}».\n"
                                "They will be marked «Sheet missing». Remove the file anyway?") != QMessageBox.Yes:
            return
        prefix = f"[{lab}]"
        for i in reversed(range(self.tabs.count())):
            w = self.tabs.widget(i)
            if str(w.property("key")).startswith(prefix):
                self.views.pop(w.property("key"), None)
                self.tabs.removeTab(i)
                w.deleteLater()
        self.book.remove_external(lab)
        self._update_srcbtn()
        self.refresh_rules()

    def _update_srcbtn(self):
        n = len(self.book.ext)
        self.btn_srcfiles.setText(f"Source files ({n}) ▾" if n else "Source files ▾")

    # ── pair by clicking: source cell, then its target cell ─────────────────
    def _pair_toggled(self, on: bool):
        self._pair_src = None
        self._pair_hint()
        self._paint()

    def _pair_reset(self):
        if self._pair_src is not None:
            self._pair_src = None
            self._pair_hint()
            self._paint()

    def _pair_hint(self):
        if not self.btn_pair.isChecked():
            self.lbl_hint.setText(HINT_NORMAL)
        elif self._pair_src is None:
            self.lbl_hint.setText("Pair mode: click a SOURCE cell, then the TARGET cell it must go to. "
                                  "Each pair is added instantly.")
        else:
            self.lbl_hint.setText(f"Source: {fmt_ref(*self._pair_src)}  →  now click its TARGET cell "
                                  "(in this file's own sheets).   Esc = cancel")

    def _cell_clicked(self, table: QTableWidget, row: int, col: int):
        if not self.btn_pair.isChecked():
            return
        key = table.property("key")
        coord = f"{get_column_letter(col + 1)}{row + 1}"
        if self._pair_src is None:
            self._pair_src = (key, coord)
            self._pair_hint()
            self._paint()
            return
        if not self.book.is_main(key):
            self.lbl_hint.setText("The target must be a cell of this file's own sheets. Click the target again.   Esc = cancel")
            return
        self.rules.append(new_rule([list(self._pair_src)], [key, coord]))
        self._pair_src = None
        self._changed()
        self._pair_hint()

    # ── selection ↔ text ────────────────────────────────────────────────────
    def take_selection(self, edit: QLineEdit, append: bool):
        key = self._cur_key()
        parts = []
        for rg in self.views[key].selectedRanges():
            a = f"{get_column_letter(rg.leftColumn() + 1)}{rg.topRow() + 1}"
            b = f"{get_column_letter(rg.rightColumn() + 1)}{rg.bottomRow() + 1}"
            parts.append(f"{qsheet(key)}!{a if a == b else a + ':' + b}")
        if not parts:
            QMessageBox.information(self, APP_NAME, "Select some cells in the sheet preview first.")
            return
        txt = "; ".join(parts)
        edit.setText(f"{edit.text().strip()}; {txt}" if append and edit.text().strip() else txt)

    def _goto(self, sheet: str, coord: str):
        t = self.views.get(sheet)
        if not t:
            return
        r, c = split_coord(coord)
        self.tabs.setCurrentWidget(t)
        t.clearSelection()
        t.setCurrentCell(r - 1, c - 1)
        t.scrollToItem(t.item(r - 1, c - 1), QAbstractItemView.PositionAtCenter)

    # ── mappings ────────────────────────────────────────────────────────────
    def add_mapping(self):
        cur = self._cur_key()
        by_col = self.cmb_order.currentData() == "col"
        try:
            src = parse_refs(self.ed_src.text(), cur, by_col)
            dst = parse_refs(self.ed_dst.text(), cur, by_col)
        except ValueError as e:
            QMessageBox.warning(self, APP_NAME, str(e))
            return
        if not src or not dst:
            QMessageBox.information(self, APP_NAME, "Fill both Source and Target.")
            return
        for sh, _ in src:
            if not self.book.has(sh):
                QMessageBox.warning(self, APP_NAME, f"Sheet not found: «{key_label(sh)}».\n"
                                                    "For another file, attach it first with «Source files».")
                return
        for sh, _ in dst:
            if not self.book.is_main(sh):
                QMessageBox.warning(self, APP_NAME, f"Target cells must be in this file's own sheets (not «{key_label(sh)}»).")
                return
        if len(src) == len(dst):
            new = [new_rule([s], d) for s, d in zip(src, dst)]          # every source cell ↔ exactly one target cell
        elif len(dst) == 1:
            new = [new_rule(src, dst[0], self.cmb_mode.currentData())]
        else:
            QMessageBox.warning(self, APP_NAME,
                                f"Source has {len(src)} cells but Target has {len(dst)}.\n\n"
                                "Every source cell must correspond to exactly one target cell: select the same number "
                                "of cells on both sides (or a single target cell to combine all sources into it).")
            return
        if self.chk_preview.isChecked() and (len(new) > 1 or len(new[0]["src"]) > 1):
            if PairPreviewDialog(self.book, new, self).exec() != QDialog.Accepted:
                return
        self.rules += new
        self.ed_src.clear()
        self.ed_dst.clear()
        self._changed()

    def _changed(self):
        self.dirty = True
        self.refresh_rules()

    def _rule_at(self, row: int) -> Optional[dict]:
        it = self.tbl.item(row, COL_CHK)
        if not it:
            return None
        rid = it.data(Qt.UserRole)
        return next((r for r in self.rules if r["id"] == rid), None)

    def _value(self, r: dict) -> Any:
        return rule_value(self.book, r, self.overrides.get(r["id"]), self.chk_digits.isChecked())

    def _status(self, r: dict, dups: Counter) -> tuple[str, bool]:
        return rule_status(self.book, r, self._value(r), r["id"] in self.overrides, dups, self.chk_over.isChecked())

    def refresh_rules(self):
        self._loading = True
        dups = Counter(tuple(r["dst"]) for r in self.rules if r["enabled"])
        self.tbl.setRowCount(len(self.rules))
        will = 0
        for i, r in enumerate(self.rules):
            chk = QTableWidgetItem()
            chk.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            chk.setCheckState(Qt.Checked if r["enabled"] else Qt.Unchecked)
            chk.setData(Qt.UserRole, r["id"])
            self.tbl.setItem(i, COL_CHK, chk)
            self.tbl.setItem(i, COL_N, _ro(str(i + 1), True))
            s_txt = "; ".join(fmt_ref(*s) for s in r["src"])
            s_it = _ro(s_txt)
            s_it.setToolTip(s_txt)
            self.tbl.setItem(i, COL_SRC, s_it)
            self.tbl.setItem(i, COL_DST, _ro(fmt_ref(*r["dst"])))
            self.tbl.setItem(i, COL_MODE, _ro(MODES[r["mode"]].split(" (")[0] if len(r["src"]) > 1 else "—", True))
            vi = QTableWidgetItem(fmt_display(self._value(r)))
            vi.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsEditable)
            if self.book.is_main(r["dst"][0]):
                ex = self.book.get(*r["dst"])
                if not is_blank(ex):
                    vi.setToolTip(f"Current value in target: {fmt_display(ex)}")
            self.tbl.setItem(i, COL_VAL, vi)
            st, ok = self._status(r, dups)
            si = _ro(st, True)
            if st in ("Sheet missing", "Duplicate target"):
                si.setBackground(QBrush(QColor(THEME["bad"])))
            elif st in ("Overwrites", "Source empty", "Target filled"):
                si.setBackground(QBrush(QColor(THEME["warn"])))
            elif ok:
                si.setBackground(QBrush(QColor(THEME["ok"])))
            self.tbl.setItem(i, COL_STAT, si)
            will += ok
        self.tbl.resizeRowsToContents()
        self._loading = False
        self.lbl_sum.setText(f"{len(self.rules)} mappings · {will} will be written")
        self.btn_run.setEnabled(will > 0)
        self._update_tpl_label()
        self._paint()

    def _paint(self):
        for sh, r, c in self._painted:
            it = self.views[sh].item(r, c) if sh in self.views else None
            if it:
                it.setBackground(QBrush())
        self._painted = set()
        jobs = []
        for rule in self.rules:
            jobs += [(THEME["src"], rule["src"]), (THEME["dst"], [rule["dst"]])]
        if self._pair_src is not None:
            jobs.append((THEME["src"], [list(self._pair_src)]))
        for color, cells in jobs:
            for sh, co in cells:
                t = self.views.get(sh)
                if not t:
                    continue
                r, c = split_coord(co)
                r, c = r - 1, c - 1
                if r >= t.rowCount() or c >= t.columnCount():
                    continue
                it = t.item(r, c)
                if it is None:
                    it = QTableWidgetItem("")
                    it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
                    t.setItem(r, c, it)
                it.setBackground(QBrush(QColor(color)))
                self._painted.add((sh, r, c))

    def _item_changed(self, item: QTableWidgetItem):
        if self._loading:
            return
        r = self._rule_at(item.row())
        if not r:
            return
        if item.column() == COL_CHK:
            r["enabled"] = item.checkState() == Qt.Checked
            self._changed()
        elif item.column() == COL_VAL:
            txt = item.text().strip()
            auto = fmt_display(rule_value(self.book, r, None, self.chk_digits.isChecked()))
            if txt and txt != auto:
                self.overrides[r["id"]] = txt
            else:
                self.overrides.pop(r["id"], None)
            self.refresh_rules()

    def _dbl(self, row: int, col: int):
        r = self._rule_at(row)
        if not r:
            return
        if col in (COL_SRC, COL_DST):
            cur = "; ".join(fmt_ref(*s) for s in r["src"]) if col == COL_SRC else fmt_ref(*r["dst"])
            text, ok = QInputDialog.getText(self, "Edit cells", "Cells (e.g.  Survey!B5 ; Info!C3):",
                                            QLineEdit.Normal, cur)
            if not ok:
                return
            try:
                cells = parse_refs(text, r["dst"][0] if col == COL_DST else r["src"][0][0])
                if not cells or (col == COL_DST and len(cells) != 1):
                    raise ValueError("Target must be exactly one cell; source needs at least one.")
                for sh, _ in cells:
                    if not self.book.has(sh):
                        raise ValueError(f"Sheet not found: «{key_label(sh)}»")
                    if col == COL_DST and not self.book.is_main(sh):
                        raise ValueError("The target must be in this file's own sheets.")
            except ValueError as e:
                QMessageBox.warning(self, APP_NAME, str(e))
                return
            if col == COL_SRC:
                r["src"] = [list(c) for c in cells]
            else:
                r["dst"] = list(cells[0])
            self._changed()
        elif col == COL_MODE and len(r["src"]) > 1:
            names = list(MODES.values())
            pick, ok = QInputDialog.getItem(self, "Mode", "When several sources feed this target:", names,
                                            list(MODES).index(r["mode"]), False)
            if ok:
                r["mode"] = list(MODES)[names.index(pick)]
                self._changed()

    def _ctx_menu(self, pos):
        row = self.tbl.rowAt(pos.y())
        r = self._rule_at(row) if row >= 0 else None
        if not r:
            return
        m = QMenu(self)
        a1 = m.addAction("Go to source cell")
        a2 = m.addAction("Go to target cell")
        a3 = m.addAction("Remove my manual value")
        a3.setEnabled(r["id"] in self.overrides)
        m.addSeparator()
        a4 = m.addAction("Delete this mapping")
        act = m.exec(self.tbl.viewport().mapToGlobal(pos))
        if act == a1:
            self._goto(*r["src"][0])
        elif act == a2:
            self._goto(*r["dst"])
        elif act == a3:
            self.overrides.pop(r["id"], None)
            self.refresh_rules()
        elif act == a4:
            self.rules = [x for x in self.rules if x["id"] != r["id"]]
            self._changed()

    def move(self, d: int):
        rows = sorted({i.row() for i in self.tbl.selectedItems()})
        if len(rows) != 1:
            return
        i, j = rows[0], rows[0] + d
        if not 0 <= j < len(self.rules):
            return
        self.rules[i], self.rules[j] = self.rules[j], self.rules[i]
        self._changed()
        self.tbl.selectRow(j)

    def delete_selected(self):
        ids = {self._rule_at(i.row())["id"] for i in self.tbl.selectedItems() if self._rule_at(i.row())}
        if ids:
            self.rules = [r for r in self.rules if r["id"] not in ids]
            self._changed()

    def clear_all(self):
        if self.rules and QMessageBox.question(self, APP_NAME, "Remove all mappings from this file's list?") == QMessageBox.Yes:
            self.rules, self.overrides = [], {}
            self._changed()

    # ── templates ───────────────────────────────────────────────────────────
    def refresh_templates(self):
        self.cmb_tpl.blockSignals(True)
        self.cmb_tpl.clear()
        self.cmb_tpl.addItem("— No template —", "")
        for n in self.templates.names():
            self.cmb_tpl.addItem(n, n)
        idx = self.cmb_tpl.findData(self.cur_template)
        if idx < 0:
            self.cur_template = ""
        self.cmb_tpl.setCurrentIndex(max(idx, 0))
        self.cmb_tpl.blockSignals(False)
        self._update_tpl_label()

    def _update_tpl_label(self):
        if self.cur_template:
            self.lbl_tpl.setText(f"  Using «{self.cur_template}»" + ("  ·  modified, not saved" if self.dirty else ""))
        else:
            self.lbl_tpl.setText("  No template" + ("  ·  unsaved mappings" if self.rules else ""))

    def _set_combo(self, name: str):
        self.cmb_tpl.blockSignals(True)
        self.cmb_tpl.setCurrentIndex(max(self.cmb_tpl.findData(name), 0))
        self.cmb_tpl.blockSignals(False)

    def _tpl_chosen(self):
        name = self.cmb_tpl.currentData() or ""
        if name == self.cur_template:
            return
        if not name:
            self.cur_template = ""
            self._tpl_sheet_map = {}
            self._update_tpl_label()
            return
        if self.rules and self.dirty and QMessageBox.question(
                self, APP_NAME, "Replace the current (unsaved) mappings with this template?") != QMessageBox.Yes:
            self._set_combo(self.cur_template)
            return
        index, alias = self.templates.info(name)
        tpl_rules = self.templates.get(name)
        self._attach_template_sources(name, tpl_rules)
        res = remap_rules(self, name, self.book, tpl_rules, index, alias, self.path.name)
        if res is None:
            self._set_combo(self.cur_template)
            return
        rules, choice, smap = res
        if choice:
            try:                                    # remember the choices for the next file
                self.templates.set_aliases(name, {**alias, **{k: v for k, v in choice.items() if v}})
            except Exception:
                pass
        self.rules = rules
        self.overrides = {}
        self.seq = self.templates.get_seq(name)
        self._tpl_sheet_map = smap
        self.cur_template, self.dirty = name, False
        self.refresh_rules()
        bad = sum(1 for r in self.rules if not rule_sheets_ok(self.book, r))
        if bad:
            QMessageBox.warning(self, APP_NAME, f"{bad} mapping(s) refer to sheets that don't exist in this file "
                                                "(marked «Sheet missing»). Edit them, attach the source file, or use a different template.")

    def _attach_template_sources(self, name: str, rules: list[dict]):
        """Attaches the extra source files a template needs: from its saved location, from the files already
        open in the app, or - if neither works - asks the user to locate them."""
        saved = self.templates.ext_files(name)
        needed = list(dict.fromkeys(Book.split_key(s[0])[0] for r in rules for s in r["src"] if s[0].startswith("[")))
        for lab in needed:
            if lab in self.book.ext:
                continue
            p = Path(saved[lab]) if saved.get(lab) else None
            if p and p.exists() and p.resolve() != self.path.resolve():
                self.add_source_file(p)
                continue
            opened = next((o.path for o in self.others_provider() if o.path.name == lab), None)
            if opened:
                self.add_source_file(opened)
                continue
            if QMessageBox.question(self, APP_NAME, f"This template reads from «{lab}», which is not attached and "
                                                    "was not found at its saved location.\n\nLocate it now?") == QMessageBox.Yes:
                fn, _ = QFileDialog.getOpenFileName(self, f"Locate {lab}", str(self.path.parent), "Excel (*.xlsx *.xlsm)")
                if fn and Path(fn).resolve() != self.path.resolve():
                    self.add_source_file(Path(fn))
                    try:
                        self.templates.set_ext_file(name, lab, fn)       # remember the new location
                    except Exception:
                        pass

    def _persist(self, name: str):
        same = name == self.cur_template
        # keep the template's ORIGINAL sheet names, so it still works with the files it was made for
        rev = {v: k for k, v in self._tpl_sheet_map.items()} if same else {}
        rules = deepcopy(self.rules)
        for r in rules:
            r["src"] = [[rev.get(s[0], s[0]), s[1]] for s in r["src"]]
            r["dst"] = [rev.get(r["dst"][0], r["dst"][0]), r["dst"][1]]
        index = {}
        for r in rules:
            for sh in [s[0] for s in r["src"]] + [r["dst"][0]]:
                cur = self._tpl_sheet_map.get(sh, sh) if same else sh
                if not sh.startswith("[") and cur in self.book.main_sheets():
                    index[sh] = self.book.main_sheets().index(cur)
        ext_files = {}
        for r in rules:
            for sh in [s[0] for s in r["src"]]:
                if sh.startswith("["):
                    cur = self._tpl_sheet_map.get(sh, sh) if same else sh
                    lab_cur, lab_tpl = Book.split_key(cur)[0], Book.split_key(sh)[0]
                    if lab_cur in self.book.ext_paths:
                        ext_files[lab_tpl] = str(self.book.ext_paths[lab_cur])
        try:
            self.templates.put(name, rules, index, ext_files, self.seq)
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, f"Could not save the template:\n{ex}")
            return
        if not same:
            self._tpl_sheet_map = {}
        self.cur_template, self.dirty = name, False
        self.templates_changed.emit()
        self.refresh_templates()

    def save_template(self):
        if not self.rules:
            QMessageBox.information(self, APP_NAME, "Add some mappings first.")
        elif self.cur_template:
            self._persist(self.cur_template)
        else:
            self.save_template_as()

    def save_template_as(self):
        if not self.rules:
            QMessageBox.information(self, APP_NAME, "Add some mappings first.")
            return
        name, ok = QInputDialog.getText(self, "Save template", "Template name:", QLineEdit.Normal, self.cur_template)
        name = name.strip()
        if not ok or not name:
            return
        if name in self.templates.data and name != self.cur_template and QMessageBox.question(
                self, APP_NAME, f"A template named «{name}» already exists. Replace it?") != QMessageBox.Yes:
            return
        self._persist(name)

    def delete_template(self):
        name = self.cmb_tpl.currentData()
        if name and QMessageBox.question(self, APP_NAME, f"Delete template «{name}»? (Your current list stays.)") == QMessageBox.Yes:
            try:
                self.templates.delete(name)
            except Exception as ex:
                QMessageBox.warning(self, APP_NAME, str(ex))
                return
            self.cur_template = ""
            self._tpl_sheet_map = {}
            self.templates_changed.emit()
            self.refresh_templates()

    # ── sequential fill ─────────────────────────────────────────────────────
    def open_sequential(self):
        labels = list(dict.fromkeys(Book.split_key(s[0])[0] for r in self.rules if r["enabled"]
                                    for s in r["src"] if s[0].startswith("[")))
        labels = [l for l in labels if l in self.book.ext]
        if not labels:
            QMessageBox.information(
                self, APP_NAME, "How to use Sequential fill:\n\n"
                                "1) Attach ONE sample source file («Source files ▾»).\n"
                                "2) Map its cells to the FIRST record of the table you want to fill "
                                "(e.g. Survey!B3 → Sheet1!C2).\n"
                                "3) Press «Sequential fill…», add all the source files, preview, and execute.")
            return
        SequentialDialog(self, labels).exec()

    # ── export (this file, all open files, or some of them) ─────────────────
    def _write_file(self, path: Path, writes: list, out: Path) -> tuple[Path, int]:
        try:
            return write_output(path, writes, out)
        except PermissionError:
            fn, _ = QFileDialog.getSaveFileName(self, f"Folder not writable — choose where to save ({path.name})",
                                                str(Path.home() / out.name), "Excel (*.xlsx *.xlsm)")
            if not fn:
                raise _Cancelled()
            return write_output(path, writes, Path(fn))

    def _pick_files(self, others: list) -> Optional[list]:
        dlg = FilePickDialog(others, self)
        if dlg.exec() != QDialog.Accepted:
            return None
        return dlg.chosen()

    def _run_on(self, t: "FilePanel") -> tuple:
        """Applies THIS panel's mappings to another open file without touching that panel's own state."""
        name = t.path.name
        try:
            book = Book(t.path)
            for p in self.book.ext_paths.values():          # same extra source files
                try:
                    book.add_external(p)
                except Exception:
                    pass
            pos = {n: i for i, n in enumerate(self.book.main_sheets())}      # sheet positions here -> hint for the other file
            res = remap_rules(self, self.cur_template or "the current mappings", book, deepcopy(self.rules),
                              index=pos, file_label=name)
            if res is None:
                return (name, None, None, "skipped (sheet matching cancelled)")
            writes = plan_writes(book, res[0], self.chk_over.isChecked(), self.chk_digits.isChecked())
            if not writes:
                return (name, None, 0, "nothing to write")
            out, n = self._write_file(t.path, writes, unique_output_path(t.path))
            return (name, out, n, "")
        except _Cancelled:
            return (name, None, None, "cancelled")
        except Exception as ex:
            return (name, None, None, f"error: {ex}")

    def run_export(self):
        dups = Counter(tuple(r["dst"]) for r in self.rules if r["enabled"])
        writes = [(r["dst"][0], r["dst"][1], self._value(r)) for r in self.rules if self._status(r, dups)[1]]
        if not writes:
            QMessageBox.information(self, APP_NAME, "Nothing to write.")
            return
        out = unique_output_path(self.path)
        others = list(self.others_provider())
        box = QMessageBox(self)
        box.setWindowTitle("Confirm")
        box.setIcon(QMessageBox.Question)
        box.setText(f"<b>{len(writes)}</b> cells will be filled in «{self.path.name}»; "
                    f"{len(self.rules) - len(writes)} mappings skipped.")
        b_all = b_some = None
        if others:
            box.setInformativeText(f"Apply the same operation to the {len(others)} other open file(s) too?\n"
                                   "Every result is saved as a new copy next to its original.")
            b_all = box.addButton("Yes, all open files", QMessageBox.AcceptRole)
            b_some = box.addButton("Some files…", QMessageBox.AcceptRole)
            b_one = box.addButton("No, only this file", QMessageBox.AcceptRole)
            box.addButton("Cancel", QMessageBox.RejectRole)
        else:
            box.setInformativeText(f"The new file will be saved as:\n{out.name}\nnext to the original, which stays unchanged.")
            b_one = box.addButton("Execute", QMessageBox.AcceptRole)
            box.addButton("Back", QMessageBox.RejectRole)
        box.exec()
        clicked = box.clickedButton()
        if b_all is not None and clicked is b_all:
            targets = others
        elif b_some is not None and clicked is b_some:
            targets = self._pick_files(others)
            if targets is None:
                return
        elif clicked is b_one:
            targets = []
        else:
            return

        results = []
        try:
            out, n = self._write_file(self.path, writes, out)
            results.append((self.path.name, out, n, ""))
        except _Cancelled:
            return
        except Exception:
            QMessageBox.critical(self, APP_NAME, f"Save failed:\n\n{traceback.format_exc()[-1200:]}")
            return
        for t in targets:
            results.append(self._run_on(t))

        done = QMessageBox(self)
        done.setWindowTitle("Done")
        done.setIcon(QMessageBox.Information)
        if len(results) == 1:
            done.setText(f"<b>{results[0][2]}</b> cells written. File saved:\n{results[0][1]}")
        else:
            total = sum(r[2] or 0 for r in results)
            okc = sum(1 for r in results if r[1])
            done.setText(f"<b>{total}</b> cells written in {okc} of {len(results)} files.")
            done.setInformativeText("\n".join(
                f"✓ {r[1].name} — {r[2]} cells" if r[1] else f"✗ {r[0]} — {r[3]}" for r in results))
        op = done.addButton("Open folder", QMessageBox.ActionRole)
        done.addButton("Close", QMessageBox.RejectRole)
        done.exec()
        if done.clickedButton() is op:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(out.parent)))


class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_NAME)
        scr = QApplication.primaryScreen().availableGeometry()
        self.setMinimumSize(min(900, scr.width()), min(560, scr.height()))
        self.resize(min(1440, int(scr.width() * 0.92)), min(860, int(scr.height() * 0.9)))
        self.move(scr.center() - self.rect().center())
        self.templates = Templates()
        self.panels: dict[str, FilePanel] = {}
        self.setAcceptDrops(True)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(12, 12, 12, 4)
        outer.setSpacing(4)
        root = QHBoxLayout()
        root.setSpacing(12)
        outer.addLayout(root, 1)

        side = card()
        side.setFixedWidth(220)
        sl = QVBoxLayout(side)
        sl.setContentsMargins(14, 14, 14, 14)
        t = QLabel("Files")
        t.setObjectName("title")
        sl.addWidget(t)
        b_add = QPushButton("＋  Add Excel files")
        b_add.setObjectName("primary")
        b_add.clicked.connect(self.add_files)
        sl.addWidget(b_add)
        self.lst = QListWidget()
        self.lst.currentRowChanged.connect(self._select)
        sl.addWidget(self.lst, 1)
        b_rm = QPushButton("Remove selected file")
        b_rm.clicked.connect(self.remove_current)
        sl.addWidget(b_rm)
        self.btn_theme = QPushButton("")
        self.btn_theme.clicked.connect(self.toggle_theme)
        sl.addWidget(self.btn_theme)
        self.lbl_tpl = QLabel("")
        self.lbl_tpl.setObjectName("muted")
        self.lbl_tpl.setWordWrap(True)
        sl.addWidget(self.lbl_tpl)
        self._tpl_label()
        root.addWidget(side)

        self.stack = QStackedWidget()
        self.welcome = QLabel("Drag Excel files here\nor click «Add Excel files»\n\nEach file is processed completely independently")
        self.welcome.setAlignment(Qt.AlignCenter)
        self.welcome.setObjectName("muted")
        self.stack.addWidget(self.welcome)
        root.addWidget(self.stack, 1)

        sig = QPushButton("Developed by Dev. Marcel Soufi")
        sig.setObjectName("sig")
        sig.setCursor(Qt.PointingHandCursor)
        sig.setFlat(True)
        sig.clicked.connect(self.contact_dev)
        outer.addWidget(sig, 0, Qt.AlignHCenter)
        self._theme_ui()

    # ── theme ───────────────────────────────────────────────────────────────
    def _theme_ui(self):
        self.btn_theme.setText("☀  Light mode" if THEME_NAME == "dark" else "🌙  Dark mode")
        self.welcome.setStyleSheet(f"font-size: 15pt; border: 2px dashed {THEME['border2']}; "
                                   f"border-radius: 16px; background: {THEME['card']};")

    def toggle_theme(self):
        new = "light" if THEME_NAME == "dark" else "dark"
        apply_theme(QApplication.instance(), new)
        save_settings({**load_settings(), "theme": new})
        self._theme_ui()
        for p in self.panels.values():
            p.on_theme_changed()

    def contact_dev(self):
        if QMessageBox.question(self, APP_NAME, "Would you like to contact the developer?") == QMessageBox.Yes:
            QDesktopServices.openUrl(QUrl(f"mailto:{DEV_EMAIL}?subject=Excel%20Report%20Filler"))

    def _tpl_label(self):
        self.lbl_tpl.setText(f"Saved templates: {len(self.templates.names())}")

    def _templates_changed(self):
        self._tpl_label()
        for p in self.panels.values():
            p.refresh_templates()

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        self.open_paths([u.toLocalFile() for u in e.mimeData().urls()])

    def add_files(self):
        fns, _ = QFileDialog.getOpenFileNames(self, "Choose Excel files", "", "Excel (*.xlsx *.xlsm)")
        self.open_paths(fns)

    def open_paths(self, paths: list[str]):
        for p in paths:
            pp = Path(p)
            if pp.suffix.lower() not in (".xlsx", ".xlsm"):
                continue
            key = str(pp.resolve())
            if key in self.panels:
                self.lst.setCurrentRow(list(self.panels).index(key))
                continue
            try:
                panel = FilePanel(pp, self.templates)
                panel.templates_changed.connect(self._templates_changed)
                panel.others_provider = lambda me=panel: [x for x in self.panels.values() if x is not me]
            except Exception as ex:
                QMessageBox.warning(self, APP_NAME, f"Could not open {pp.name}:\n{ex}")
                continue
            self.panels[key] = panel
            self.stack.addWidget(panel)
            it = QListWidgetItem(f"📄 {pp.name}")
            it.setToolTip(key)
            self.lst.addItem(it)
            self.lst.setCurrentRow(self.lst.count() - 1)

    def _select(self, i: int):
        if i < 0:
            self.stack.setCurrentWidget(self.welcome)
            return
        self.stack.setCurrentWidget(list(self.panels.values())[i])

    def remove_current(self):
        i = self.lst.currentRow()
        if i < 0:
            return
        w = self.panels.pop(list(self.panels)[i])
        self.stack.removeWidget(w)
        w.deleteLater()
        self.lst.takeItem(i)


def main():
    if sys.platform == "win32":                    # show our own icon in the taskbar (not python's)
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("excel.report.filler")
        except Exception:
            pass
    app = QApplication(sys.argv)
    if ICON_FILE.exists():
        app.setWindowIcon(QIcon(str(ICON_FILE)))
    app.setStyle("Fusion")
    app.setLayoutDirection(Qt.LeftToRight)
    apply_theme(app, default_theme())
    w = MainWindow()
    w.show()
    for a in sys.argv[1:]:
        w.open_paths([a])
    sys.exit(app.exec())


if __name__ == "__main__":
    main()